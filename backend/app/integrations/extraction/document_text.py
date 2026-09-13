"""Deterministic, LLM-free text extraction — no network calls, ever.

Used two ways:

1. Building the LLM input for formats Claude cannot ingest as a `document`/
   `image` content block (XLSX, CSV, DOCX) — Claude reads plain text
   natively, so there's no reason to render a spreadsheet as an image or
   run it through vision.
2. Independently verifying an LLM's claimed `source_quote` actually
   appears in the document, for every format where a reliable text layer
   exists — see `claude_extraction.py`'s `verify_quote`.

Returns `None` when a format/file has no reliable text layer (a scanned
PDF with no text layer at all, or a raster image) — extraction still
proceeds by sending the raw bytes to Claude's vision, just without
independent quote verification, which is exactly why those cases are
capped at MEDIUM confidence rather than HIGH (see `resolve_confidence`).
"""

import io


def extract_pages(*, file_bytes: bytes, extension: str) -> list[str] | None:
    """One string per page (PDF), sheet (XLSX), or the whole document (CSV,
    DOCX — formats with no natural page boundary become a single "page 1").
    `None` for formats with no text layer at all (png/jpg/jpeg), or a PDF
    that turns out to have no extractable text (scanned).
    """
    if extension == "pdf":
        return _extract_pdf_pages(file_bytes)
    if extension == "xlsx":
        return _extract_xlsx_sheets(file_bytes)
    if extension == "csv":
        return [file_bytes.decode("utf-8", errors="replace")]
    if extension == "docx":
        return _extract_docx_text(file_bytes)
    return None


def _extract_pdf_pages(file_bytes: bytes) -> list[str] | None:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception:
        return None
    if not any(page.strip() for page in pages):
        return None  # scanned PDF, no text layer to verify against
    return pages


def _extract_xlsx_sheets(file_bytes: bytes) -> list[str] | None:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        sheets = []
        for sheet in workbook.worksheets:
            lines = [
                ",".join("" if cell is None else str(cell) for cell in row)
                for row in sheet.iter_rows(values_only=True)
            ]
            sheets.append("\n".join(lines))
    except Exception:
        return None
    return sheets or [""]


def _extract_docx_text(file_bytes: bytes) -> list[str] | None:
    import docx

    try:
        document = docx.Document(io.BytesIO(file_bytes))
        return ["\n".join(paragraph.text for paragraph in document.paragraphs)]
    except Exception:
        return None


def normalize(text: str) -> str:
    """Whitespace/case normalization shared by extraction and verification
    so a quote spanning a line break, or differing only in case, still
    matches."""
    return " ".join(text.split()).lower()
